import csv, json
from datetime import datetime
from pathlib import Path
from .database import Database

HEADERS=["Company","Role","Location","Status","Applied Date","Deadline","Interview Date",
         "Salary Min","Salary Max","Contact","Job URL","Notes"]

def _rows(db,search="",status="All"):
    for r in db.list(search,status):
        yield [r["company"],r["role"],r["location"],r["status"],r["applied_date"] or "",
               r["deadline"] or "",r["interview_date"] or "",r["salary_min"] or "",
               r["salary_max"] or "",r["contact"],r["job_url"],r["notes"]]

def export_csv(db,path,search="",status="All"):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",newline="",encoding="utf-8") as f:
        w=csv.writer(f); w.writerow(HEADERS); w.writerows(_rows(db,search,status))
    return path

def export_excel(db,path,search="",status="All"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); ws=wb.active; ws.title="Applications"; ws.append(HEADERS)
    for cell in ws[1]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="DCE6F1")
    for row in _rows(db,search,status): ws.append(row)
    for i,w in enumerate([24,24,18,15,14,14,17,14,14,24,42,50],1):
        ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"; wb.save(path); return path

def export_json(db,path):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    data={"generated_at":datetime.now().isoformat(timespec="seconds"),
          "statistics":db.stats(),"upcoming":[dict(r) for r in db.upcoming(30)]}
    path.write_text(json.dumps(data,indent=2,default=str),encoding="utf-8"); return path

def export_pdf(db,path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    styles=getSampleStyleSheet(); s=db.stats()
    elements=[Paragraph("Job Application Tracker — Report",styles["Title"]),
              Paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M} · Total: {s['total']} · Active: {s['active']} · Interviews/Offers: {s['interviews']}",styles["Normal"]),Spacer(1,12)]
    data=[HEADERS[:10]]+[row[:10] for row in _rows(db)]
    table=Table(data,repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#263238")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("GRID",(0,0),(-1,-1),0.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP")]))
    elements.append(table); doc.build(elements); return path
